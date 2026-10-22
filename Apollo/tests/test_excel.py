import csv
import json
from datetime import datetime

from openpyxl import Workbook, load_workbook

from apollo_tool.excel_io import read_input, write_output
from apollo_tool.models import CompanyFacts, CompanyResult, Person

HEADERS = [
    "company_name",
    "website",
    "domain",
    "team_size",
    "niche",
    "wellfound_profile_url",
    "date_scraped",
    "wellfound_company_id",
]

VERIFIED_HEADERS = [
    "Startup Name",
    "Official Website",
    "Employee Range",
    "Exact Member Count",
    "Wellfound Profile URL",
    "Website Verification Status",
    "Discovery Date",
    "Last Checked",
    "Source of Website Evidence",
    "Processing Status",
]


def make_input(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Companies"
    ws.append(HEADERS)
    ws.append([
        "Linear",
        "https://linear.app",
        "linear.app",
        "",
        "",
        "https://wellfound.com/company/linear",
        datetime(2026, 10, 1),
        "111",
    ])
    ws.append([
        "Example Co",
        "example.com",
        "example.com",
        "11-50",
        "SaaS",
        "https://wellfound.com/company/example",
        datetime(2026, 10, 2),
        "222",
    ])
    ws.append([
        "Linear duplicate",
        "https://linear.app",
        "linear.app",
        "300",
        "Productivity",
        "https://wellfound.com/company/linear",
        datetime(2026, 10, 3),
        "111",
    ])
    wb.save(path)


def test_read_input(tmp_path):
    path = tmp_path / "input.xlsx"
    make_input(path)
    wb = read_input(path)
    assert wb.sheet_name == "Companies"
    assert len(wb.rows) == 3
    assert wb.rows[0].domain == "linear.app"
    assert wb.rows[1].wellfound_company_id == "222"


def test_read_input_accepts_human_readable_headers(tmp_path):
    path = tmp_path / "input.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([
        "Company Name",
        "Website",
        "Domain",
        "Team Size",
        "Niche",
        "Wellfound Profile URL",
        "Date Scraped",
        "Wellfound Company ID",
    ])
    ws.append([
        "Linear",
        "https://linear.app",
        "linear.app",
        "11-50",
        "SaaS",
        "https://wellfound.com/company/linear",
        datetime(2026, 10, 1),
        "111",
    ])
    wb.save(path)

    source = read_input(path)
    assert source.headers[0] == "Company Name"
    assert source.rows[0].domain == "linear.app"
    assert source.rows[0].data["company_name"] == "Linear"
    assert source.rows[0].data["team_size"] == "11-50"
    assert source.rows[0].data["wellfound_company_id"] == "111"


def test_read_input_skips_cover_sheet_and_blank_column(tmp_path):
    path = tmp_path / "input.xlsx"
    wb = Workbook()
    cover = wb.active
    cover.title = "Read me"
    cover["A1"] = "Instructions"
    ws = wb.create_sheet("Data")
    # A blank spacer column between company_name and website must not shift values.
    ws["A1"] = "company_name"
    ws["C1"] = "website"
    ws["D1"] = "domain"
    ws["E1"] = "team_size"
    ws["F1"] = "niche"
    ws["G1"] = "wellfound_profile_url"
    ws["H1"] = "date_scraped"
    ws["I1"] = "wellfound_company_id"
    ws["A2"] = "Linear"
    ws["C2"] = "https://linear.app"
    ws["D2"] = "linear.app"
    ws["I2"] = "111"
    wb.save(path)

    source = read_input(path)
    assert source.sheet_name == "Data"
    assert source.headers[0:2] == ["company_name", "website"]
    assert source.rows[0].values[0:2] == ["Linear", "https://linear.app"]
    assert source.rows[0].domain == "linear.app"
    assert source.rows[0].data["wellfound_company_id"] == "111"


def make_verified_input(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Verified"
    ws.append(VERIFIED_HEADERS)
    ws.append([
        "African Minds",
        "http://africanminds.co.za",
        "1 employee",
        None,
        None,
        "verified",
        "2026-10-20T14:23:04+00:00",
        "2026-10-20T15:19:02+00:00",
        "dbpedia - HTTP 200",
        "verified",
    ])
    ws.append([
        "AIKOCorp",
        "https://aikocorp.ai",
        "1-10 Employees",
        None,
        "https://wellfound.com/company/aikocorp",
        "verified",
        "2026-10-20T18:13:52+00:00",
        "2026-10-21T03:50:22+00:00",
        "Wellfound company profile - HTTP 200",
        "verified",
    ])
    wb.save(path)


def test_read_verified_export_xlsx(tmp_path):
    path = tmp_path / "verified.xlsx"
    make_verified_input(path)

    source = read_input(path)
    assert source.headers == VERIFIED_HEADERS
    assert len(source.rows) == 2
    assert source.rows[0].company_name == "African Minds"
    assert source.rows[0].domain == "africanminds.co.za"
    assert source.rows[0].data["team_size"] == "1 employee"
    assert source.rows[1].domain == "aikocorp.ai"
    assert source.rows[1].data["wellfound_profile_url"] == "https://wellfound.com/company/aikocorp"
    assert source.rows[1].wellfound_company_id == "aikocorp"


def test_read_verified_export_csv_and_json(tmp_path):
    csv_path = tmp_path / "verified.csv"
    with csv_path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(VERIFIED_HEADERS)
        writer.writerow([
            "African Minds",
            "http://africanminds.co.za",
            "1 employee",
            "",
            "",
            "verified",
            "2026-10-20T14:23:04+00:00",
            "2026-10-20T15:19:02+00:00",
            "dbpedia - HTTP 200",
            "verified",
        ])
    csv_source = read_input(csv_path)
    assert len(csv_source.rows) == 1
    assert csv_source.rows[0].domain == "africanminds.co.za"
    assert csv_source.headers == VERIFIED_HEADERS

    json_path = tmp_path / "verified.json"
    record = dict(zip(VERIFIED_HEADERS, [
        "AIKOCorp",
        "https://aikocorp.ai",
        "1-10 Employees",
        None,
        "https://wellfound.com/company/aikocorp",
        "verified",
        "2026-10-20T18:13:52+00:00",
        "2026-10-21T03:50:22+00:00",
        "Wellfound company profile - HTTP 200",
        "verified",
    ]))
    json_path.write_text(json.dumps({"records": [record]}), encoding="utf-8")
    json_source = read_input(json_path)
    assert len(json_source.rows) == 1
    assert json_source.rows[0].domain == "aikocorp.ai"
    assert json_source.rows[0].wellfound_company_id == "aikocorp"
    assert json_source.headers == VERIFIED_HEADERS


def test_write_output_preserves_verified_columns(tmp_path):
    path = tmp_path / "verified.xlsx"
    make_verified_input(path)
    source = read_input(path)
    result = CompanyResult(
        domain="africanminds.co.za",
        company_name="African Minds",
        wellfound_company_id="",
        status="processed",
        email="hello@africanminds.co.za",
        email_source="verified",
        notes="pattern hello (from founder)",
        facts=CompanyFacts(employees="42", industries=["Education"]),
    )
    output = write_output(source, [result], tmp_path / "output.xlsx")
    wb = load_workbook(output)
    ws = wb["Verified"]
    headers = [cell.value for cell in ws[1]]
    assert headers == VERIFIED_HEADERS + ["email", "email_source", "notes"]
    rows = [
        {headers[i]: ws.cell(row=row, column=i + 1).value for i in range(len(headers))}
        for row in (2, 3)
    ]
    assert rows[0]["Employee Range"] == "1 employee"
    assert rows[0]["email"] == "hello@africanminds.co.za"
    assert rows[1]["Employee Range"] == "1-10 Employees"
    assert rows[1]["email"] is None


def test_write_output_fills_blanks_and_formats(tmp_path):
    input_path = tmp_path / "input.xlsx"
    make_input(input_path)
    source = read_input(input_path)

    linear = CompanyResult(
        domain="linear.app",
        company_name="Linear",
        wellfound_company_id="111",
        status="processed",
        email="tuomas.artman@linear.app",
        email_source="verified",
        notes="pattern first.last (from Tuomas Artman); team 3: 1 guessed, 1 masked",
        revealed_name="Tuomas Artman",
        pattern="first.last",
        credit_spent=1,
        facts=CompanyFacts(
            employees="300",
            industries=["Information Technology & Services"],
            stage="Series B",
            revenue="$100M",
        ),
        people=[
            Person("p1", "Tuomas Artman", "Tuomas Artman", False, "Co-Founder", 0,
                   email="tuomas.artman@linear.app", source="verified",
                   validation="real email revealed via Apollo", confidence="high"),
            Person("p2", "Doug Parker", "Doug Parker", False, "Head of Talent", 3,
                   email="doug.parker@linear.app", source="guessed",
                   validation="syntax_ok; mx_ok", confidence="high"),
            Person("p3", "Jamie F ------", "Jamie F", True, "Head of Security", 4,
                   email="", source="",
                   validation="skipped: last name masked by Apollo Free plan", confidence="n/a"),
        ],
    )
    example = CompanyResult(
        domain="example.com",
        company_name="Example Co",
        wellfound_company_id="222",
        status="no_candidate",
        notes="no candidate; no credit spent",
        facts=CompanyFacts(employees="20", industries=["Fintech"]),
        people=[],
    )

    output = write_output(
        source,
        [linear, example],
        tmp_path / "output.xlsx",
        statuses={},
    )
    wb = load_workbook(output)
    ws = wb["Companies"]
    headers = [cell.value for cell in ws[1]]
    assert headers[-3:] == ["email", "email_source", "notes"]
    assert ws.freeze_panes == "A2"
    assert ws["A1"].font.bold is True
    assert ws.auto_filter.ref

    row1 = {headers[i]: ws.cell(row=2, column=i + 1).value for i in range(len(headers))}
    row2 = {headers[i]: ws.cell(row=3, column=i + 1).value for i in range(len(headers))}
    row3 = {headers[i]: ws.cell(row=4, column=i + 1).value for i in range(len(headers))}

    assert row1["email"] == "tuomas.artman@linear.app"
    assert row1["email_source"] == "verified"
    assert row1["team_size"] == "300"
    assert row1["niche"] == "Information Technology & Services"
    assert row2["team_size"] == "11-50"
    assert row2["niche"] == "SaaS"
    assert row3["email"] == "tuomas.artman@linear.app"
    assert row1["wellfound_company_id"] == "111"

    team = wb["team_emails"]
    team_headers = [cell.value for cell in team[1]]
    assert team_headers == [
        "company", "domain", "full_name", "title", "email",
        "source", "syntax/mx validation", "confidence",
    ]
    assert team.max_row == 4
    assert team["E2"].value == "tuomas.artman@linear.app"
    assert team["F2"].value == "verified"

    assert wb.sheetnames[0] == "employees"
    employees = wb["employees"]
    employee_headers = [cell.value for cell in employees[1]]
    assert employee_headers == [
        "Employee Name",
        "Company Name",
        "Company Members",
        "Website",
        "Job Title",
        "Email",
    ]
    assert employee_headers[-1] == "Email"
    # Example Co has no collected contacts, so it is intentionally absent from
    # the people-first sheet (it still exists in the original-data sheet).
    assert employees.max_row == 4
    assert employees["A2"].value == "Tuomas Artman"
    assert employees["B2"].value == "Linear"
    assert employees["C2"].value == "300"
    assert employees["D2"].value == "https://linear.app"
    assert employees["F2"].value == "tuomas.artman@linear.app"
    assert employees["F2"].fill.fgColor.rgb.endswith("C6EFCE")
    assert employees["F3"].fill.fgColor.rgb.endswith("FFEB9C")
    assert employees["F4"].value is None
    assert employees["F4"].fill.fgColor.rgb.endswith("E7E6E6")
    assert employees["F4"].comment is not None
    assert "masked" in employees["F4"].comment.text.lower()
    assert employees["F1"].comment is not None
    assert "Green = verified" in employees["F1"].comment.text

    assert "Run summary" in wb.sheetnames
    summary = wb["Run summary"]
    summary_values = [
        [cell.value for cell in row]
        for row in summary.iter_rows()
    ]
    flat = "\n".join(
        " | ".join(str(value) for value in row if value is not None)
        for row in summary_values
    )
    assert "Live run" in flat
    assert "Contacts listed" in flat
    assert "2 (67%)" in flat
    assert "Pattern guesses (high confidence)" in flat
    assert "1" in flat
    assert "Credits spent" in flat
